import streamlit as st
import json
from pdfplumber import PDF
import docx
from fpdf import FPDF
from datetime import datetime

st.set_page_config(page_title="Whole-Brain Tender Questions", page_icon="🧠", layout="wide")

# ====================== LOGO ======================
col1, col2 = st.columns([1, 4])

with col1:
    st.image("VS logo 2.png", width=180)   # Change width as needed

with col2:
    st.title("Whole-Brain Tender Analyser")
    st.markdown("This tool is designed to be used in conjunction with your HBDI profile to help you get the most out of whole brain thinking when writing a tender by leaning into your non dominant areas")
    st.markdown("Instructions: Upload a tender/RFP - Click the big red button.")   
# ====================== MODEL SELECTOR ======================
model_options = {
    "Gemini 2.5 Flash (Google - FREE to start)": "gemini",
    "Grok 4.1 Fast (xAI - cheap & powerful)": "grok"
}
selected_model_name = st.selectbox("Choose AI Model:", options=list(model_options.keys()), index=0)
selected_model = model_options[selected_model_name]

# ====================== API KEYS ======================
if selected_model == "grok":
    api_key = st.secrets.get("XAI_API_KEY")
    if not api_key:
        st.warning("⚠️ Add your XAI_API_KEY in Streamlit secrets to use Grok.")
else:
    api_key = st.secrets.get("GEMINI_API_KEY")
    if not api_key:
        st.warning("⚠️ Add your GEMINI_API_KEY in Streamlit secrets to use Gemini.")

# ====================== FILE UPLOAD ======================
uploaded_file = st.file_uploader("Upload tender document (PDF or Word)", type=["pdf", "docx"])

def extract_text(file):
    if file.type == "application/pdf":
        text = ""
        with PDF(file) as pdf:
            for page in pdf.pages:
                text += (page.extract_text() or "") + "\n"
        return text
    elif file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        doc = docx.Document(file)
        return "\n".join([para.text for para in doc.paragraphs])
    return ""

# ====================== PDF GENERATION FUNCTION (Tighter Layout + Logo) ======================
def create_whole_brain_pdf(questions):
    class PDF(FPDF):
        def header(self):
            # Add logo on every page
            if hasattr(self, 'logo_path') and self.logo_path:
                self.image(self.logo_path, x=10, y=0, w=45)  # ← Adjust size/position here
            
            # Title centered
            self.set_font("Helvetica", "B", 16)
            self.cell(0, 10, "Whole-Brain Tender Approach", ln=1, align="C")
            
            self.set_font("Helvetica", "", 10)
            self.cell(0, 6, f"Generated on {datetime.now().strftime('%d %B %Y')}", ln=1, align="C")
            self.ln(12)

        def quadrant_title(self, title, rgb):
            self.set_font("Helvetica", "B", 14)
            self.set_fill_color(*rgb)
            self.set_text_color(255, 255, 255)
            self.cell(0, 12, title, ln=1, align="C", fill=True)
            self.ln(8)

        def check_and_new_page(self, lines_needed=3):
            if self.get_y() > self.h - 35:   # Leave margin at bottom
                self.add_page()

    pdf = PDF()
    pdf.logo_path = "VS logo 2.png"           # Make sure your logo is named exactly this
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()

    colors = {
        "A": (0, 102, 204),
        "B": (0, 153, 0),
        "C": (204, 0, 0),
        "D": (204, 153, 0)
    }

    quadrant_names = {
        "A": "Quadrant A - Analytical (Blue)",
        "B": "Quadrant B - Practical (Green)",
        "C": "Quadrant C - Relational (Red)",
        "D": "Quadrant D - Conceptual (Yellow)"
    }

    for q in ["A", "B", "C", "D"]:
        pdf.quadrant_title(quadrant_names[q], colors[q])
        
        pdf.set_text_color(0, 0, 0)
        pdf.set_font("Helvetica", "", 10)   # Smaller font for fitting 10 questions
        
        for i, question in enumerate(questions.get(q, []), 1):
            clean_question = question.replace("–", "-").replace("—", "-").replace("’", "'")
            
            pdf.check_and_new_page()
            
            pdf.multi_cell(0, 7, f"{i}. {clean_question}")
            pdf.ln(6)

        # New page for next quadrant (except after the last one)
        if q != "D":
            pdf.add_page()

    pdf_output = pdf.output(dest="S")
    if isinstance(pdf_output, bytearray):
        return bytes(pdf_output)
    return pdf_output

# ====================== MAIN LOGIC ======================
if uploaded_file:
    with st.spinner("Extracting text from document..."):
        tender_text = extract_text(uploaded_file)
    
    if len(tender_text.strip()) < 200:
        st.error("Not enough text extracted. Please try a different file.")
        st.stop()

    st.success(f"✅ Extracted ~{len(tender_text):,} characters")

    if st.button("Generate my Whole-Brain Tender check list", type="primary", use_container_width=True):
        if not api_key:
            st.error("Missing API key for the selected model. Please check Streamlit secrets.")
            st.stop()

        with st.spinner(f"Analysing tender with {selected_model_name}..."):
            prompt = """You are an expert bid strategist and HBDI specialist.
Analyse this tender document and create exactly 10 specific, actionable questions for EACH HBDI quadrant.

Tender text:
=== TENDER START ===
""" + tender_text + """
=== TENDER END ===

Output ONLY valid JSON with this exact structure:
{
  "A": ["question 1", "question 2", ...],
  "B": ["question 1", ...],
  "C": ["question 1", ...],
  "D": ["question 1", ...]
}

Quadrants:
- A (Blue): Analytical, data, ROI, technical specs, risk, metrics
- B (Green): Process, timelines, compliance, steps, organisation
- C (Red): People, stakeholders, culture, communication, values, team fit
- D (Yellow): Innovation, big picture, future vision, creative alternatives"""

            try:
                if selected_model == "grok":
                    from openai import OpenAI
                    client = OpenAI(api_key=api_key, base_url="https://api.x.ai/v1")
                    response = client.chat.completions.create(
                        model="grok-4-1-fast-reasoning",
                        messages=[{"role": "user", "content": prompt}],
                        temperature=0.3,
                        response_format={"type": "json_object"}
                    )
                    result = json.loads(response.choices[0].message.content)
                else:  # gemini
                    import google.generativeai as genai
                    genai.configure(api_key=api_key)
                    model = genai.GenerativeModel('gemini-2.5-flash')
                    response = model.generate_content(
                        prompt,
                        generation_config={"response_mime_type": "application/json", "temperature": 0.3}
                    )
                    result = json.loads(response.text)

                st.session_state.questions = result

                st.markdown("### 🎯 Your Whole-Brain Tender Questions")

                col1, col2 = st.columns(2)
                col3, col4 = st.columns(2)

                quadrants = {
                    "A": ("🔵 Quadrant A – Analytical (Blue)", col1),
                    "B": ("🟢 Quadrant B – Practical (Green)", col2),
                    "C": ("🔴 Quadrant C – Relational (Red)", col3),
                    "D": ("🟡 Quadrant D – Conceptual (Yellow)", col4),
                }

                for q, (title, col) in quadrants.items():
                    with col:
                        st.markdown(f"**{title}**")
                        for i, question in enumerate(result.get(q, []), 1):
                            st.markdown(f"{i}. {question}")

            except Exception as e:
                st.error(f"Error generating questions: {str(e)}")

# ====================== PDF DOWNLOAD ======================
if "questions" in st.session_state:
    st.markdown("---")
    if st.button("📄 Download PDF", type="primary", use_container_width=True):
        with st.spinner("Creating beautiful PDF..."):
            try:
                pdf_bytes = create_whole_brain_pdf(st.session_state.questions)
                st.download_button(
                    label="⬇️ Click here to download the PDF",
                    data=pdf_bytes,
                    file_name=f"Whole_Brain_Tender_Questions_{datetime.now().strftime('%Y-%m-%d')}.pdf",
                    mime="application/pdf",
                    key="pdf_download"
                )
            except Exception as pdf_error:
                st.error(f"PDF creation failed: {str(pdf_error)}")

st.caption("Whole-Brain Tender Tool • Value Shift")
